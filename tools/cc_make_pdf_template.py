# Build the app's form-fillable PDF sheet from the PowerPoint template.
#
#   1. python tools/cc_make_pdf_template.py stage1 <blank.pptx> <spec.json>
#        writes a copy of the template with every tag-bearing box emptied, plus a spec
#        describing where each box sits and what text belongs in it.
#   2. (PowerPoint converts <blank.pptx> to a PDF - see cc_build_pdf.ps1)
#   3. python tools/cc_make_pdf_template.py stage2 <blank.pdf> <spec.json>
#        drops an AcroForm field onto each of those boxes and writes
#        resources/character_creator_template.pdf + resources/pdf-sheet-fields.js
#
# Each box becomes ONE field carrying its whole text, headings included ("ACTIONS\n
# {actions}\n\nBONUS ACTIONS\n{bonus_actions}"). Filling per tag instead would put a
# one-line field above a heading and let long values run straight over it.
import sys,re,io,json,copy
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

TOK=re.compile(r"<([A-Za-z_0-9]+)>")
EMU_IN=914400.0
PT_IN=72.0
TEMPLATE="resources/character_creator_template.pptx"

def all_shapes(shapes):
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            for s in all_shapes(sh.shapes): yield s
        else: yield sh

def first_run(sh):
    for p in sh.text_frame.paragraphs:
        if p.runs: return p.runs[0]
    return None

def field_name(tags,used):
    base="f_"+tags[0]
    n=base; i=2
    while n in used: n=base+"_%d"%i; i+=1
    used.add(n); return n

def stage1(blank_path,spec_path):
    prs=Presentation(TEMPLATE)
    spec={"page_w":prs.slide_width/EMU_IN*PT_IN,"page_h":prs.slide_height/EMU_IN*PT_IN,"fields":[]}
    used=set()
    for pi,slide in enumerate(prs.slides):
        for sh in all_shapes(slide.shapes):
            if not sh.has_text_frame: continue
            text="\n".join("".join(r.text for r in p.runs) for p in sh.text_frame.paragraphs)
            tags=TOK.findall(text)
            if not tags: continue
            r=first_run(sh)
            size=(r.font.size.pt if (r is not None and r.font.size) else 10.0)
            bold=bool(r.font.bold) if r is not None else False
            colour=None
            try:
                if r is not None and r.font.color and r.font.color.rgb is not None:
                    colour=str(r.font.color.rgb)
            except Exception: pass
            try: align=str(sh.text_frame.paragraphs[0].alignment or "")
            except Exception: align=""
            checkbox=len(tags)==1 and (tags[0].startswith("st_") or tags[0].startswith("chk_"))
            spec["fields"].append({
                "name":field_name(tags,used),
                "tags":tags,
                "template":TOK.sub(lambda m:"{"+m.group(1)+"}",text),
                "page":pi,
                "rect":[sh.left/EMU_IN*PT_IN,sh.top/EMU_IN*PT_IN,
                        (sh.left+sh.width)/EMU_IN*PT_IN,(sh.top+sh.height)/EMU_IN*PT_IN],
                "size":size,"bold":bold,"colour":colour,"align":align,
                "checkbox":checkbox,
                # a single-line field gets centred vertically by the viewer, which is what
                # the template's MIDDLE anchor does; only mark a field multiline when it
                # really holds several lines, or the small number boxes sit at the top
                "multiline":("\n" in text) or (sh.height/EMU_IN*PT_IN)>=size*2.6,
            })
            # empty the box: the field will carry the headings as well as the values
            p0=sh.text_frame.paragraphs[0]
            for rr in list(p0.runs): rr.text=""
            for p in list(sh.text_frame.paragraphs)[1:]: p._p.getparent().remove(p._p)
            if checkbox:
                sh.fill.background()          # the tick is drawn by the field itself
    prs.save(blank_path)
    io.open(spec_path,"w",encoding="utf-8").write(json.dumps(spec,indent=1))
    print("stage1: %d fields over %d slides -> %s"%(len(spec["fields"]),len(prs.slides.__iter__.__self__._sldIdLst) if False else len(prs.slides._sldIdLst),blank_path))

def stage2(pdf_path,spec_path):
    import fitz
    spec=json.load(io.open(spec_path,encoding="utf-8"))
    doc=fitz.open(pdf_path)
    made=0
    for f in spec["fields"]:
        if f["page"]>=doc.page_count: continue
        page=doc[f["page"]]
        # PowerPoint may scale the slide onto the page; map by ratio rather than assuming
        sx=page.rect.width/spec["page_w"]; sy=page.rect.height/spec["page_h"]
        x0,y0,x1,y1=f["rect"]
        rect=fitz.Rect(x0*sx,y0*sy,x1*sx,y1*sy)
        w=fitz.Widget()
        w.rect=rect
        w.field_name=f["name"]
        w.border_width=0
        w.fill_color=None
        if f["checkbox"]:
            w.field_type=fitz.PDF_WIDGET_TYPE_CHECKBOX
            w.field_value=False
        else:
            w.field_type=fitz.PDF_WIDGET_TYPE_TEXT
            w.field_value=""
            # Helvetica runs wider than the template's Aptos Narrow, so the boxes that hold
            # flowing text get a point back to keep the wrapping close to the PowerPoint look
            w.text_fontsize=max(6.0,f["size"]-1) if f["multiline"] else f["size"]
            w.text_font="Helv"
            if f["colour"]:
                try:
                    c=f["colour"]; w.text_color=[int(c[i:i+2],16)/255.0 for i in (0,2,4)]
                except Exception: pass
            flags=0
            if f["multiline"]: flags|=fitz.PDF_TX_FIELD_IS_MULTILINE
            w.field_flags=flags
        annot=page.add_widget(w)
        # Quadding (/Q) is how a form field carries its alignment; this build of PyMuPDF
        # has no Widget attribute for it, so write the key onto the field itself.
        q=1 if "CENTER" in f["align"] else 2 if "RIGHT" in f["align"] else 0
        f["q"]=q
        if q and annot is not None:
            try: doc.xref_set_key(annot.xref,"Q","%d"%q)
            except Exception: pass
        made+=1
    out="resources/character_creator_template.pdf"
    doc.save(out,garbage=4,deflate=True)
    print("stage2: %d widgets -> %s"%(made,out))
    # the per-field text templates the app composes from
    js={f["name"]:{"t":f["template"],"cb":f["checkbox"],"tags":f["tags"],
                   "q":f.get("q",0),"ml":f["multiline"]} for f in spec["fields"]}
    io.open("resources/pdf-sheet-fields.js","w",encoding="utf-8").write(
        "// Auto-generated from character_creator_template.pptx by tools/cc_make_pdf_template.py\n"
        "// field name -> {t: text template, cb: checkbox?, q: 0 left 1 centre 2 right, ml: multiline?}\n"
        "window.CC_SHEET_FIELDS = "+json.dumps(js,ensure_ascii=False)+";\n")
    print("        %d field templates -> resources/pdf-sheet-fields.js"%len(js))
    # the app runs from file://, where fetch() of a local PDF is blocked, so the sheet
    # travels as base64 in a script file just like the WotC one it replaces
    import base64
    b64=base64.b64encode(io.open(out,"rb").read()).decode("ascii")
    io.open("resources/pdf-sheet-template.js","w",encoding="utf-8").write(
        "// Auto-generated from character_creator_template.pptx by tools/cc_make_pdf_template.py\n"
        'window.CC_SHEET_TEMPLATE="'+b64+'";\n')
    print("        %d KB base64 -> resources/pdf-sheet-template.js"%(len(b64)//1024))

if __name__=="__main__":
    if sys.argv[1]=="stage1": stage1(sys.argv[2],sys.argv[3])
    elif sys.argv[1]=="stage2": stage2(sys.argv[2],sys.argv[3])
    else: raise SystemExit("stage1 or stage2")
