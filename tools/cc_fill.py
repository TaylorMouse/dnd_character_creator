# Fill the PowerPoint character-sheet template from a data file.
#   python tools/cc_fill.py <data.txt> <out.pptx>
# The data file holds one "tag|value" record per line (as produced by cc_fetch.js);
# a literal \n inside a value becomes a real line break in the slide.
#
# A box whose content is longer than it can show spills onto continuation pages: the
# slide is cloned, its header gets "(cont.)", and the remaining text carries over. That
# is what lets the spell list and the features & traits run to any length.
import sys,re,io,copy
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.text.text import _Paragraph

data_path,out=sys.argv[1],sys.argv[2]
DATA={}
for line in io.open(data_path,encoding="utf-8"):
    line=line.rstrip("\n").rstrip("\r")
    if "|" in line:
        k,v=line.split("|",1)
        DATA[k]=v.replace("\\n","\n")          # unescape the newlines cc_fetch encoded

tok=re.compile(r"<([A-Za-z_0-9]+)>")
DOT=RGBColor(0x33,0x33,0x33)
# Only these boxes may run onto continuation pages. Every other box is a fixed part of
# the character-sheet layout and has to stay exactly where the template puts it.
def is_flow(tag): return tag=="features_and_traits" or tag.startswith("spell_list")
EMU_PER_IN=914400.0
PT_PER_IN=72.0
# a rough character width as a fraction of the font size; narrow faces are tighter.
# Deliberately generous so a page never overflows in PowerPoint's own layout.
WIDTH_RATIO={"narrow":0.47,"normal":0.505}
LINE_SPACING=1.22

def all_shapes(shapes):
    for sh in shapes:
        if sh.shape_type==MSO_SHAPE_TYPE.GROUP:
            for s in all_shapes(sh.shapes): yield s
        else: yield sh

def first_run(shape):
    for para in shape.text_frame.paragraphs:
        if para.runs: return para.runs[0]
    return None

def capacity(shape):
    """How many characters per line and lines per page this box can hold."""
    r=first_run(shape)
    size=(r.font.size.pt if (r is not None and r.font.size) else 10.0)
    name=(r.font.name if r is not None else "") or ""
    ratio=WIDTH_RATIO["narrow"] if "narrow" in name.lower() else WIDTH_RATIO["normal"]
    tf=shape.text_frame
    li=tf.margin_left if tf.margin_left is not None else 91440
    ri=tf.margin_right if tf.margin_right is not None else 91440
    ti=tf.margin_top if tf.margin_top is not None else 45720
    bi=tf.margin_bottom if tf.margin_bottom is not None else 45720
    w_pt=max(1.0,(shape.width-li-ri)/EMU_PER_IN*PT_PER_IN)
    h_pt=max(1.0,(shape.height-ti-bi)/EMU_PER_IN*PT_PER_IN)
    cpl=max(20,int(w_pt/(size*ratio)))
    lpp=max(4,int(h_pt/(size*LINE_SPACING)))
    return cpl,lpp

def drop_empty_sections(text):
    """The spell box carries a fixed header per level ("LEVEL 7 SPELLS"). Drop the ones
    whose section turned out empty, so a low-level caster doesn't get pages of headings."""
    def is_head(l):
        s=l.strip()
        return bool(s) and s==s.upper() and any(c.isalpha() for c in s)
    lines=text.split("\n"); out=[]; i=0
    while i<len(lines):
        if not is_head(lines[i]): out.append(lines[i]); i+=1; continue
        j=i+1; body=[]
        while j<len(lines) and not is_head(lines[j]): body.append(lines[j]); j+=1
        if any(b.strip() for b in body): out.append(lines[i]); out.extend(body)
        i=j
    while out and not out[-1].strip(): out.pop()
    return "\n".join(out)

def wrap(line,width):
    """Word-wrap one logical line; returns at least one (possibly empty) display line."""
    if not line: return [""]
    words,cur,outl=line.split(" "),"",[]
    for w in words:
        if not cur: cur=w
        elif len(cur)+1+len(w)<=width: cur+=" "+w
        else:
            outl.append(cur); cur=w
        while len(cur)>width:                     # a single very long word
            outl.append(cur[:width]); cur=cur[width:]
    outl.append(cur)
    return outl

def paginate(text,cpl,lpp):
    """Split text into pages of display lines, breaking between blank-line-separated
    blocks where possible so a feature's heading stays with its text."""
    blocks=text.split("\n\n")
    pages,cur=[],[]
    for bi,blk in enumerate(blocks):
        lines=[]
        for ln in blk.split("\n"): lines.extend(wrap(ln,cpl))
        if cur: lines=[""]+lines               # blank spacer between blocks
        if len(cur)+len(lines)<=lpp:
            cur.extend(lines); continue
        if cur:                                 # start a fresh page for this block
            pages.append(cur); cur=[]
            lines=lines[1:] if lines and lines[0]=="" else lines
        while len(lines)>lpp:                   # block longer than a whole page
            pages.append(lines[:lpp]); lines=lines[lpp:]
        cur=lines
    if cur: pages.append(cur)
    for p in pages:                             # never open a page on a blank spacer
        while p and not p[0].strip(): p.pop(0)
    return pages or [[]]

def set_lines(shape,lines):
    """Write pre-wrapped lines into a shape, keeping the placeholder run's formatting."""
    tf=shape.text_frame
    first=tf.paragraphs[0]
    for r in list(first.runs)[1:]: r._r.getparent().remove(r._r)
    if first.runs: first.runs[0].text=lines[0] if lines else ""
    else: first.text=lines[0] if lines else ""
    for para in list(tf.paragraphs)[1:]: para._p.getparent().remove(para._p)
    for ln in (lines[1:] if lines else []):
        p=copy.deepcopy(first._p)
        first._p.getparent().append(p)
        np=_Paragraph(p,first._parent)
        if np.runs:
            np.runs[0].text=ln
            for r in list(np.runs)[1:]: r._r.getparent().remove(r._r)
        else: np.text=ln

RID=("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id",
     "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
     "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link")

def clone_slide(prs,src):
    """Copy a slide (shapes + the image relationships they point at) to the end."""
    new=prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes): shp._element.getparent().remove(shp._element)
    remap={}
    for rId,rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT,RT.NOTES_SLIDE): continue
        remap[rId]=(new.part.rels.get_or_add_ext_rel(rel.reltype,rel._target)
                    if rel.is_external else new.part.rels.get_or_add(rel.reltype,rel._target))
    for shp in src.shapes:
        el=copy.deepcopy(shp._element)
        for node in el.iter():                  # point copied refs at the new rIds
            for attr in RID:
                if attr in node.attrib and node.attrib[attr] in remap:
                    node.attrib[attr]=remap[node.attrib[attr]]
        new.shapes._spTree.append(el)
    return new

def delete_slide(prs,slide):
    """Drop a slide entirely - used when a flow page (the spell list of a non-caster)
    turns out to have nothing on it."""
    lst=prs.slides._sldIdLst
    rid_attr="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for e in list(lst):
        if int(e.get("id"))==slide.slide_id:
            prs.part.drop_rel(e.get(rid_attr)); lst.remove(e); return

def move_slide(prs,slide,index):
    lst=prs.slides._sldIdLst
    for e in list(lst):
        if int(e.get("id"))==slide.slide_id:
            lst.remove(e); lst.insert(index,e); return

filled=set(); missing=set()
prs=Presentation("resources/character_creator_template.pptx")

# ---- pass 1: substitute every tag on every slide ----
overflow=[]      # (slide, shape_name, pages) for boxes that need continuation slides
blank_flow=[]    # flow slides that ended up with nothing to show
for slide in prs.slides:
    for sh in all_shapes(slide.shapes):
        if not sh.has_text_frame: continue
        full="".join(r.text for para in sh.text_frame.paragraphs for r in para.runs)
        names=tok.findall(full)
        if not names: continue
        # a lone st_/chk_ tag marks a proficiency dot: fill the shape instead of writing text
        if len(names)==1 and (names[0].startswith("st_") or names[0].startswith("chk_")):
            filled.add(names[0])
            on=DATA.get(names[0],"0")=="1"
            for para in sh.text_frame.paragraphs:
                for r in para.runs: r.text=""
            if on: sh.fill.solid(); sh.fill.fore_color.rgb=DOT
            else: sh.fill.background()
            continue
        # substitute into the shape's whole text, keeping its line structure
        lines=[]
        for para in sh.text_frame.paragraphs:
            ptxt="".join(r.text for r in para.runs)
            def sub(mm):
                if mm.group(1) in DATA:
                    filled.add(mm.group(1)); return DATA[mm.group(1)]
                missing.add(mm.group(1)); return mm.group(0)
            lines.append(tok.sub(sub,ptxt))
        text="\n".join(lines)
        if not any(is_flow(n) for n in names):
            set_lines(sh,text.split("\n"))          # fixed sheet box: leave the layout alone
            continue
        if any(n.startswith("spell_list") for n in names): text=drop_empty_sections(text)
        if not text.strip(): blank_flow.append(slide); continue   # e.g. a non-caster's spell page
        cpl,lpp=capacity(sh)
        pages=paginate(text,cpl,lpp)
        set_lines(sh,pages[0])
        if len(pages)>1: overflow.append((slide,sh.name,pages))

# ---- pass 2: spill long boxes onto continuation slides ----
added=0
for slide,shape_name,pages in overflow:
    base=list(prs.slides).index(slide)
    for n,page in enumerate(pages[1:],start=1):
        clone=clone_slide(prs,slide)
        for sh in all_shapes(clone.shapes):
            if sh.name==shape_name: set_lines(sh,page)
            elif sh.has_text_frame:
                t=sh.text_frame.text
                if t and t.isupper() and "(CONT." not in t:      # the box header
                    r=first_run(sh)
                    if r is not None: r.text=t+" (CONT. %d)"%(n+1)
        move_slide(prs,clone,base+n)
        added+=1

for slide in blank_flow: delete_slide(prs,slide)

prs.save(out)
print("filled -> %s  (%d tags, %d page%s%s)"%(out,len(filled),len(prs.slides.__iter__.__self__._sldIdLst),
      "" if len(prs.slides._sldIdLst)==1 else "s",
      ", +%d continuation"%added if added else ""))
if missing: print("   MISSING:", ", ".join(sorted(missing)))
